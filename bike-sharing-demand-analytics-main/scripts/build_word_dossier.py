"""Generate Microsoft Word (.doc and .docx) versions of the HR & Tech Submission Dossier
with all charts, tables, and executive formatting embedded."""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src import config

SOURCE_MD = config.REPORTS_DIR / "HR_and_Tech_Submission_Dossier.md"
SOURCE_HTML = config.REPORTS_DIR / "HR_and_Tech_Submission_Dossier.html"

DOCX_REPORT = config.REPORTS_DIR / "HR_and_Tech_Submission_Dossier.docx"
DOC_REPORT = config.REPORTS_DIR / "HR_and_Tech_Submission_Dossier.doc"
DOCX_ROOT = ROOT / "HR_and_Tech_Submission_Dossier.docx"
DOC_ROOT = ROOT / "HR_and_Tech_Submission_Dossier.doc"


def _shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _clean_text(text: str) -> str:
    # Convert markdown file links [label](file://...) -> label
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    # Clean basic LaTeX inline symbols for Word readability
    replacements = {
        r"\mu_{\text{working}}": "μ_working",
        r"\mu_{\text{non-working}}": "μ_non-working",
        r"\alpha": "α",
        r"\implies": "⇒",
        r"\times": "×",
        r"\le": "≤",
        r"\ge": "≥",
        r"\to": "→",
        r"\in": "∈",
        r"^\circ\text{C}": "°C",
        r"R^2": "R²",
        r"H_0": "H₀",
        r"H_1": "H₁",
        r"H_2": "H₂",
        r"z_{\text{robust}}": "z_robust",
        r"\text{df}": "df",
        r"\text{SD}": "SD",
        r"\text{Durbin-Watson}": "Durbin-Watson",
        r"\log(1 + y) / \exp(\hat{y})-1": "log(1 + y) / exp(ŷ) - 1",
        r"\log(1 + y) / \exp(\hat{y}) - 1": "log(1 + y) / exp(ŷ) - 1",
        r"\log(1+y)": "log(1 + y)",
        r"\log(1 + \text{cnt})": "log(1 + cnt)",
        r"e^{\hat{\beta}_1}": "exp(β̂₁)",
        r"\%": "%",
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    return text


def _add_formatted_runs(paragraph, text: str) -> None:
    text = _clean_text(text)
    # Tokenize by **bold** and `code`
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        else:
            paragraph.add_run(tok)


def _flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            val = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _add_formatted_runs(p, val)
            for run in p.runs:
                run.font.size = Pt(9.0)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            if r_idx == 0:
                _shade_cell(cell, "E2E8F0")
            elif r_idx % 2 == 1:
                _shade_cell(cell, "F8FAFC")
    doc.add_paragraph()


def build_docx() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    table_buffer: list[list[str]] = []
    in_code = False
    code_lines: list[str] = []

    for line in lines:
        # Handle fenced code / mermaid blocks
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                if code_lines and not any("flowchart " in cl for cl in code_lines[:2]):
                    p = doc.add_paragraph()
                    r = p.add_run("\n".join(code_lines))
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.0)
                else:
                    # Readable pipeline summary for Mermaid block in Word
                    p = doc.add_paragraph()
                    r = p.add_run(
                        "Pipeline Flow: [1. Official UCI Data (hour.csv 17,379 rows, day.csv 731 rows)] "
                        "→ [2. Validation & Feature Pipeline (src/data_validation.py, src/features.py)] "
                        "→ [3. Statistical & ML Core (src/hypothesis.py, src/train_model.py, src/anomaly.py)] "
                        "→ [4. Operations Prototype (src/query_engine.py, src/predict.py, app.py)]"
                    )
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            continue
        if in_code:
            code_lines.append(line)
            continue

        # Handle tables
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                continue
            table_buffer.append(cells)
            continue
        elif table_buffer:
            _flush_table(doc, table_buffer)
            table_buffer = []

        # Skip LaTeX display math lines or format cleanly
        if line.strip().startswith("$$") and line.strip().endswith("$$"):
            eq = _clean_text(line.strip().strip("$"))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(eq)
            r.italic = True
            r.font.name = "Cambria Math"
            continue

        # Handle embedded images
        img_match = re.match(r"!\[([^\]]*)\]\(\.\./outputs/figures/([^)]+\.png)\)", line.strip())
        if img_match:
            caption, fname = img_match.group(1), img_match.group(2)
            img_path = config.FIGURES_DIR / fname
            if img_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(str(img_path), width=Inches(6.4))
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rc = p_cap.add_run(f"Figure: {caption}")
                rc.italic = True
                rc.font.size = Pt(9.0)
                rc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            continue

        if line.startswith("# "):
            h = doc.add_heading(_clean_text(line[2:]), level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        elif line.startswith("## "):
            h = doc.add_heading(_clean_text(line[3:]), level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif line.startswith("### "):
            h = doc.add_heading(_clean_text(line[4:]), level=2)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        elif line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs(p, line.strip()[2:])
        elif re.match(r"^\d+\.\s+", line.strip()):
            p = doc.add_paragraph(style="List Number")
            item_txt = re.sub(r"^\d+\.\s+", "", line.strip())
            _add_formatted_runs(p, item_txt)
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _add_formatted_runs(p, line[1:].strip())
            for r in p.runs:
                r.italic = True
        elif line.strip() == "---":
            continue
        elif line.strip():
            p = doc.add_paragraph()
            _add_formatted_runs(p, line.strip())

    if table_buffer:
        _flush_table(doc, table_buffer)

    doc.save(str(DOCX_REPORT))
    shutil.copy2(DOCX_REPORT, DOCX_ROOT)


def build_word_doc() -> None:
    """Create a Microsoft Word .doc file with Word Office XML headers and embedded images."""
    raw_html = SOURCE_HTML.read_text(encoding="utf-8")
    word_html = raw_html.replace(
        '<html lang="en">',
        '<html xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:w="urn:schemas-microsoft-com:office:word" '
        'xmlns="http://www.w3.org/TR/REC-html40">',
    ).replace(
        "<head>",
        "<head><!--[if gte mso 9]><xml><w:WordDocument><w:View>Print</w:View>"
        "<w:Zoom>100</w:Zoom><w:DoNotOptimizeForBrowser/></w:WordDocument></xml><![endif]-->",
    )
    DOC_REPORT.write_text(word_html, encoding="utf-8")
    shutil.copy2(DOC_REPORT, DOC_ROOT)


if __name__ == "__main__":
    build_docx()
    build_word_doc()
    print(f"Created DOCX: {DOCX_ROOT} ({DOCX_ROOT.stat().st_size:,} bytes)")
    print(f"Created DOC:  {DOC_ROOT} ({DOC_ROOT.stat().st_size:,} bytes)")

